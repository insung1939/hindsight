"""docs/script.md → docs/script.docx (발표 대본, 읽기 쉬운 Word 문서)."""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
FONT = "Apple SD Gothic Neo"; MONO = "Menlo"
NAVY = RGBColor(0x1B, 0x2A, 0x4A); BLUE = RGBColor(0x1E, 0x5E, 0xFF); GRAY = RGBColor(0x6B, 0x72, 0x80); RED = RGBColor(0xC0, 0x26, 0x2B)

doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2)

def set_font(run, size=None, bold=None, color=None, italic=None, mono=False):
    name = MONO if mono else FONT
    run.font.name = name; r = run._element.get_or_add_rPr(); rf = r.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); r.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"): rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color

st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.45

def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor); tcPr.append(s)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]|\*[^*\s][^*]*?\*)")
def add_inline(p, text, base_size=12.5, base_color=None):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); set_font(r, base_size, True, BLUE)
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); set_font(r, base_size - 1.5, None, NAVY, mono=True)
        elif part.startswith("[") and part.endswith("]"):
            r = p.add_run(part); set_font(r, base_size - 1, None, GRAY, italic=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); set_font(r, base_size, None, base_color, italic=True)
        else:
            r = p.add_run(part); set_font(r, base_size, None, base_color)

def heading(text, level):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(22 if level == 1 else 12); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if level == 1:
        # 장 번호와 시간을 분리해 표시: "4. 문제 → 답 (0:55)"
        m = re.match(r"^(\d+|부록 \d+)\.\s*(.+?)\s*(\(.+\))?$", text)
        if m:
            r = p.add_run(m.group(1) + "  "); set_font(r, 20, True, BLUE)
            r = p.add_run(m.group(2)); set_font(r, 20, True, NAVY)
            if m.group(3): r = p.add_run("   " + m.group(3)); set_font(r, 13, False, GRAY)
        else:
            r = p.add_run(text); set_font(r, 20, True, NAVY)
        pPr = p._element.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "D0D5DD")): bt.set(qn(k), v)
        b.append(bt); pPr.append(b)
    else:
        r = p.add_run(text); set_font(r, 15, True, NAVY)

def table(rows):
    head, body = rows[0], rows[2:] if rows[1].startswith("|--") or set(rows[1].replace("|", "").strip()) <= set("-: ") else rows[1:]
    cells = lambda ln: [c.strip() for c in ln.strip().strip("|").split("|")]
    h = cells(head); data = [cells(x) for x in body]
    t = doc.add_table(rows=1 + len(data), cols=len(h)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(h):
        cell = t.rows[0].cells[j]; shade(cell, "1B2A4A"); p = cell.paragraphs[0]; r = p.add_run(c); set_font(r, 11, True, RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(data, 1):
        for j, c in enumerate(row[:len(h)]):
            cell = t.rows[i].cells[j]; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.2
            if i % 2 == 0: shade(cell, "F3F5F9")
            add_inline(p, c, 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0; first_rule = True
while i < len(lines):
    ln = lines[i]
    if ln.startswith("# "):
        p = doc.add_paragraph(); r = p.add_run(ln[2:].replace(" — ", "\n")); set_font(r, 26, True, NAVY); p.paragraph_format.space_after = Pt(4)
        i += 1; continue
    if ln.startswith("## "): heading(ln[3:], 1); i += 1; continue
    if ln.startswith("### "): heading(ln[4:], 2); i += 1; continue
    if ln.strip() == "---":
        if not first_rule: doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first_rule = False; i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"): rows.append(lines[i]); i += 1
        table(rows); continue
    if re.match(r"^\s*- ", ln):
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.4
        add_inline(p, re.sub(r"^\s*- ", "", ln)); i += 1; continue
    if re.match(r"^\d+\. ", ln):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)
        n, rest = ln.split(". ", 1); r = p.add_run(n + ".  "); set_font(r, 12.5, True, BLUE); add_inline(p, rest); i += 1; continue
    if not ln.strip(): i += 1; continue
    # 문단: 붙어 있는 줄은 하나의 문단으로
    buf = [ln]
    while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||- |\d+\. |---)", lines[i + 1]): i += 1; buf.append(lines[i])
    p = doc.add_paragraph()
    for k, b in enumerate(buf):
        if k: p.add_run().add_break()
        add_inline(p, b)
    i += 1

# 머리글·바닥글
hp = sec.header.paragraphs[0]; r = hp.add_run("HINDSIGHT 발표 대본 · 5조 · 2026-10-21 · 비공개"); set_font(r, 9, None, GRAY); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(); set_font(r, 9, None, GRAY)
for tag, txt in (("w:fldChar", "begin"), ("w:instrText", " PAGE "), ("w:fldChar", "end")):
    el = OxmlElement(tag)
    if tag == "w:fldChar": el.set(qn("w:fldCharType"), txt)
    else: el.set(qn("xml:space"), "preserve"); el.text = txt
    r._element.append(el)
doc.save(OUT); print("wrote", OUT, OUT.stat().st_size, "bytes")
